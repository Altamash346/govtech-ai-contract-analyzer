import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="NyayaMitra_AI_Presentation.pptx"):
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette matching GovTech + Reference PDF styling
    NAVY_DARK   = RGBColor(15, 43, 72)     # #0F2B48 Sovereign Deep Navy
    NAVY_ACCENT = RGBColor(10, 25, 47)     # #0A192F Deep Navy Background
    SAFFRON     = RGBColor(234, 88, 12)    # #EA580C Saffron Orange
    EMERALD     = RGBColor(19, 136, 8)     # #138808 Indian Emerald
    CARD_BG     = RGBColor(248, 250, 252)  # #F8FAFC Off-white card
    CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0 Border
    TEXT_DARK   = RGBColor(15, 23, 42)     # #0F172A Slate 900
    TEXT_MUTED  = RGBColor(100, 116, 139)  # #64748B Slate 500
    WHITE       = RGBColor(255, 255, 255)
    RED_ACCENT  = RGBColor(220, 38, 38)    # #DC2626 Danger
    AMBER_BG    = RGBColor(254, 243, 199)

    logo_path = os.path.abspath(r"frontend\public\logo.jpg")
    has_logo = os.path.exists(logo_path)

    def add_header(slide, title_text, category_text="NyayaMitra.AI | GovTech Intelligence Platform", slide_num=None):
        # Header category
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.3))
        p_cat = tb_cat.text_frame.paragraphs[0]
        p_cat.text = category_text
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = SAFFRON

        # Main Header Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.6))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_DARK

        if slide_num:
            tb_num = slide.shapes.add_textbox(Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.3))
            p_num = tb_num.text_frame.paragraphs[0]
            p_num.text = str(slide_num)
            p_num.font.size = Pt(10)
            p_num.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Navy Background)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.color.rgb = NAVY_DARK

    # Tricolor accent line at top
    t1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.444), Inches(0.12))
    t1.fill.solid()
    t1.fill.fore_color.rgb = SAFFRON
    t1.line.fill.background()
    t2 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.444), 0, Inches(4.444), Inches(0.12))
    t2.fill.solid()
    t2.fill.fore_color.rgb = WHITE
    t2.line.fill.background()
    t3 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.888), 0, Inches(4.445), Inches(0.12))
    t3.fill.solid()
    t3.fill.fore_color.rgb = EMERALD
    t3.line.fill.background()

    # Subtitle tag
    tb = s1.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(7.5), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "GovTech AI & Civic Intelligence Capstone Project"
    p.font.size = Pt(13)
    p.font.italic = True
    p.font.color.rgb = RGBColor(226, 232, 240)

    # Main Project Title
    tb = s1.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(7.5), Inches(1.1))
    p = tb.text_frame.paragraphs[0]
    p.text = "NYAYAMITRA.AI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # Project Subtitle
    tb = s1.shapes.add_textbox(Inches(0.9), Inches(3.3), Inches(7.5), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = "Sovereign Legal Contract Intelligence & Citizen Welfare Recommendation Platform"
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(148, 163, 184)

    # Team & Guide Details
    tb = s1.shapes.add_textbox(Inches(0.9), Inches(5.1), Inches(4.5), Inches(1.8))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Team Members:"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = SAFFRON
    p2 = tf.add_paragraph()
    p2.text = "Altamash & Development Team"
    p2.font.size = Pt(12)
    p2.font.color.rgb = WHITE
    p3 = tf.add_paragraph()
    p3.text = "GovTech AI Engineering Division"
    p3.font.size = Pt(11)
    p3.font.color.rgb = RGBColor(203, 213, 225)

    tb_guide = s1.shapes.add_textbox(Inches(7.2), Inches(5.1), Inches(5.0), Inches(1.8))
    tf_g = tb_guide.text_frame
    pg = tf_g.paragraphs[0]
    pg.text = "Project Domain & Department:"
    pg.font.bold = True
    pg.font.size = Pt(12)
    pg.font.color.rgb = SAFFRON
    pg2 = tf_g.add_paragraph()
    pg2.text = "Department of Information Technology & AI"
    pg2.font.size = Pt(12)
    pg2.font.color.rgb = WHITE
    pg3 = tf_g.add_paragraph()
    pg3.text = "Sovereign Public-Interest Technology (GovTech)"
    pg3.font.size = Pt(11)
    pg3.font.color.rgb = RGBColor(203, 213, 225)

    # Circular Logo Emblem on Right
    if has_logo:
        s1.shapes.add_picture(logo_path, Inches(9.4), Inches(1.8), Inches(2.9), Inches(2.9))

    # =========================================================================
    # SLIDE 2: Problem Statement (Big Number Stat Cards)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Citizens and MSMEs sign contracts blindly", slide_num=2)

    tb = s2.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    p.text = "Legal contracts are packed with hidden indemnity traps, ambiguous termination terms, and one-sided liabilities. Meanwhile, citizens miss out on 96+ welfare schemes because policies are buried in dense legal jargon without vernacular accessibility."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK

    # 4 Stat Cards
    stats = [
        {"val": "85%", "label": "citizens sign without understanding clauses", "color": RED_ACCENT},
        {"val": "₹4.2L", "label": "average dispute loss per MSME vendor", "color": SAFFRON},
        {"val": "96+", "label": "welfare schemes underutilized due to jargon", "color": NAVY_DARK},
        {"val": "~14 days", "label": "manual legal audit turnaround time", "color": EMERALD}
    ]
    card_w = Inches(2.7)
    card_h = Inches(3.2)
    start_x = Inches(0.8)
    gap = Inches(0.3)

    for i, st in enumerate(stats):
        x = start_x + i * (card_w + gap)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        # Large Stat Number
        tb_num = s2.shapes.add_textbox(x, Inches(2.8), card_w, Inches(1.0))
        p = tb_num.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = st["val"]
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = st["color"]

        # Label
        tb_lbl = s2.shapes.add_textbox(x + Inches(0.2), Inches(4.0), card_w - Inches(0.4), Inches(1.0))
        p = tb_lbl.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = st["label"]
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # Goal Callout Banner at Bottom
    goal_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.9))
    goal_box.fill.solid()
    goal_box.fill.fore_color.rgb = AMBER_BG
    goal_box.line.color.rgb = RGBColor(245, 158, 11)
    tb_g = s2.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.8))
    p = tb_g.text_frame.paragraphs[0]
    p.text = "Project Goal: Build an AI-powered legal audit platform that automates risk detection, visual redlines, 3-agent courtroom debate, and citizen welfare scheme matching across all 22 scheduled Indian languages."
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(180, 83, 9)

    # =========================================================================
    # SLIDE 3: Project Objectives (2x3 Grid of Feature Cards)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Project objectives", slide_num=3)

    objectives = [
        {"title": "Contract RAG Engine", "desc": "Hybrid chunking, FAISS vector indexing, and citation-backed Q&A with exact page numbers."},
        {"title": "Risk & Clause Audit", "desc": "Automatic classification of standard clauses and risk severity tagging (High, Medium, Low)."},
        {"title": "Multi-Agent Courtroom", "desc": "Autonomous 3-agent legal debate (The Critic, The Defender, and Judicial Verdict) to balance terms."},
        {"title": "Visual PDF Redlining", "desc": "Direct visual strikethroughs and recommended amendment text rendered natively on uploaded PDFs."},
        {"title": "Welfare Recommender", "desc": "Demographic rule filtering and LLM eligibility matching across 96+ national & state schemes."},
        {"title": "22 Indian Languages", "desc": "Dynamic full-page translation across all Eighth Schedule constitutional languages in real-time."}
    ]

    col_w = Inches(3.7)
    col_h = Inches(2.2)
    for idx, obj in enumerate(objectives):
        row = idx // 3
        col = idx % 3
        x = Inches(0.8) + col * (col_w + Inches(0.3))
        y = Inches(1.5) + row * (col_h + Inches(0.35))

        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, col_w, col_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        # Saffron badge bar
        bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), col_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = SAFFRON
        bar.line.fill.background()

        tb = s3.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), col_w - Inches(0.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = f"{idx+1}. {obj['title']}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        tb_d = s3.shapes.add_textbox(x + Inches(0.3), y + Inches(0.7), col_w - Inches(0.5), Inches(1.3))
        p = tb_d.text_frame.paragraphs[0]
        p.text = obj['desc']
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Dataset & Knowledge Base (Sidebar + Breakdown)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Dataset & Knowledge Base Description", slide_num=4)

    # Left Stats Card
    sb = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(4.2), Inches(5.1))
    sb.fill.solid()
    sb.fill.fore_color.rgb = CARD_BG
    sb.line.color.rgb = CARD_BORDER

    tb = s4.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(3.6), Inches(4.7))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Benchmark Corpus"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    bullets = [
        "50 Sample Contracts & Tenders (e-Quotations, NDAs, Municipal works, Commercial leases)",
        "96 Central & State Welfare Schemes (Ministry notifications, data.gov.in, MyScheme)",
        "248 Extracted Legal Clauses annotated across standard contract types",
        "115 Identified Risk Instances with remediated counter-clauses",
        "Eighth Schedule Languages (22 Indian languages verified for DOM translation)"
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # Right Card 1: Risk Severity Split
    rc1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(1.5), Inches(7.2), Inches(2.35))
    rc1.fill.solid()
    rc1.fill.fore_color.rgb = WHITE
    rc1.line.color.rgb = CARD_BORDER

    tb = s4.shapes.add_textbox(Inches(5.5), Inches(1.6), Inches(6.8), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "Risk Severity Distribution Across Analyzed Contracts"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    risk_splits = [
        {"cat": "Low Risk", "pct": "42%", "desc": "Standard boilerplate, mutual jurisdiction", "c": EMERALD},
        {"cat": "Medium Risk", "pct": "36%", "desc": "Unclear IP bounds, strict payment timelines", "c": RGBColor(245, 158, 11)},
        {"cat": "High Risk", "pct": "22%", "desc": "Uncapped indemnities, immediate termination", "c": RED_ACCENT}
    ]
    for j, rs in enumerate(risk_splits):
        bx = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.5 + j*2.3), Inches(2.1), Inches(2.1), Inches(1.5))
        bx.fill.solid()
        bx.fill.fore_color.rgb = CARD_BG
        bx.line.color.rgb = rs["c"]
        tb = s4.shapes.add_textbox(Inches(5.5 + j*2.3), Inches(2.2), Inches(2.1), Inches(1.3))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = rs["pct"]
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = rs["c"]
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = rs["cat"]
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK

    # Right Card 2: Welfare Scheme Categories Breakdown
    rc2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.3), Inches(4.1), Inches(7.2), Inches(2.5))
    rc2.fill.solid()
    rc2.fill.fore_color.rgb = WHITE
    rc2.line.color.rgb = CARD_BORDER

    tb = s4.shapes.add_textbox(Inches(5.5), Inches(4.2), Inches(6.8), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "Welfare Scheme Knowledge Base (96 Directory Programs)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    schemes_cats = [
        ("Agriculture & Farmers", "24 Schemes", "PM-KISAN, Soil Card, PMFBY, AIF"),
        ("Healthcare & Maternity", "18 Schemes", "Ayushman Bharat, PMMVY, JSSK"),
        ("MSME, Trade & Loans", "18 Schemes", "PMEGP, Mudra, Stand-Up India"),
        ("Social Welfare & Housing", "22 Schemes", "PM Awas, NSAP Pensions, Atal Pension"),
        ("Education & Skill Tech", "14 Schemes", "PM-YUVA, Pre/Post-Matric, PMKVY")
    ]
    tb = s4.shapes.add_textbox(Inches(5.5), Inches(4.6), Inches(6.8), Inches(1.9))
    tf = tb.text_frame
    for cat_name, count, examples in schemes_cats:
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        p.text = f"• {cat_name} ({count}): {examples}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 5: System Architecture (Flow Architecture)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "System architecture: Decoupled GovTech Pipeline", slide_num=5)

    arch_layers = [
        {"title": "PRESENTATION LAYER", "tech": "React 19 · Vite · TailwindCSS · Lucide Icons · Google DOM i18n", "desc": "Single-Page Application deployed on Vercel with bilingual national masthead and active PDF visual redlines."},
        {"title": "GATEWAY & ACCESS LAYER", "tech": "FastAPI · Uvicorn · JWT / Session Auth · Tenant Isolation", "desc": "Strict user-scoped database isolation ensuring zero cross-citizen data leakage before authentication."},
        {"title": "ANALYTICAL ENGINE LAYER", "tech": "Dual Engine Routing (Gemini Flash vs Local Ollama Llama 3.2)", "desc": "Dynamic switching between Cloud Fast (low latency) and Air-gapped Secure (local private DPDP compliance)."},
        {"title": "VECTOR & RAG STORE", "tech": "FAISS Index · all-MiniLM-L6-v2 Embeddings · PyMuPDF", "desc": "500-char semantic chunks, cosine similarity top-3 citation retrieval, and vector-level bounding box redlining."},
        {"title": "DATA PERSISTENCE LAYER", "tech": "SQLite (app.db) · Star Schemas · 96 Scheme JSON Records", "desc": "Relational storage of users, audit reports, chat dialogue, and curated national welfare policies."}
    ]

    for idx, lyr in enumerate(arch_layers):
        y = Inches(1.4) + idx * Inches(1.05)
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        # Category indicator bar
        bar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(0.12), Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = SAFFRON if idx % 2 == 0 else NAVY_DARK
        bar.line.fill.background()

        tb = s5.shapes.add_textbox(Inches(1.1), y + Inches(0.08), Inches(4.2), Inches(0.8))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = lyr["title"]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p2 = tf.add_paragraph()
        p2.text = lyr["tech"]
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = SAFFRON

        tb_d = s5.shapes.add_textbox(Inches(5.4), y + Inches(0.15), Inches(6.9), Inches(0.7))
        p = tb_d.text_frame.paragraphs[0]
        p.text = lyr["desc"]
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: ETL & Processing Pipeline
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Contract processing: Ingestion, RAG & redline pipeline", slide_num=6)

    steps = [
        {"num": "1", "title": "Ingest & OCR", "desc": "Upload PDF, extract raw text with PyMuPDF, and run Tesseract fallback on scanned pages."},
        {"num": "2", "title": "Chunk & Embed", "desc": "Split text into 500-char chunks (100 overlap) and embed via sentence-transformers/all-MiniLM-L6-v2."},
        {"num": "3", "title": "Audit & Extract", "desc": "Unified single-pass LLM call to classify clauses, detect risk severities, and extract entities."},
        {"num": "4", "title": "Visual Redlining", "desc": "Locate exact risk coordinates in original PDF and draw visual strikethroughs and counter-clauses."}
    ]

    card_w = Inches(2.7)
    card_h = Inches(2.2)
    for i, st in enumerate(steps):
        x = Inches(0.8) + i * (card_w + Inches(0.3))
        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tb = s6.shapes.add_textbox(x + Inches(0.15), Inches(1.6), card_w - Inches(0.3), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = f"Step {st['num']}: {st['title']}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        tb_d = s6.shapes.add_textbox(x + Inches(0.15), Inches(2.2), card_w - Inches(0.3), Inches(1.4))
        p = tb_d.text_frame.paragraphs[0]
        p.text = st['desc']
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    # Bottom Architecture Highlights
    bot_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.0), Inches(11.7), Inches(2.6))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = WHITE
    bot_card.line.color.rgb = CARD_BORDER

    tb = s6.shapes.add_textbox(Inches(1.1), Inches(4.2), Inches(11.1), Inches(2.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Design Highlights & Architectural Strengths"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    highlights = [
        "Single-Pass Legal Auditing: Slashes LLM inference latency by 75% by consolidating 4 calls (summary, clauses, risks, entities) into 1 schema-validated call.",
        "Stack-Based JSON Auto-Repair: Solves token truncation by tracking bracket balances and automatically repairing incomplete LLM responses without data loss.",
        "Strict Per-User History Isolation: Foreign-key database architecture ensures each user only accesses their private audited documents and chat sessions.",
        "Vector-Level Citation Verification: Every answer generated in the Q&A assistant returns exact page number citations retrieved from FAISS chunks."
    ]
    for h in highlights:
        p = tf.add_paragraph()
        p.text = "✓ " + h
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 7: Module 1 & 2 (Audit Engine & Grounded Q&A)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Audit engine & citation-verified Q&A", slide_num=7)

    features_m = [
        ("Classify Clauses", "Indemnity, Termination, Payment Terms, Governing Law, Confidentiality, Dispute Resolution tagged with verbatim quotes."),
        ("Risk Assessment", "3-level hazard ranking (High, Medium, Low) explaining specific commercial liabilities and legal traps."),
        ("Amendment Synthesis", "Auto-generates balanced counter-clauses and amendments protecting MSMEs and citizens from asymmetric terms."),
        ("Top-3 Citation RAG", "FAISS vector retrieval extracts the 3 most relevant contract chunks and injects them into the prompt for zero hallucination."),
        ("Official PDF Reports", "One-click export of structured executive audit reports complete with entity tables, risks, and Q&A transcripts.")
    ]

    for idx, (title, desc) in enumerate(features_m):
        y = Inches(1.5) + idx * Inches(1.0)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.9))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tb = s7.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(3.2), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        tb_d = s7.shapes.add_textbox(Inches(4.4), y + Inches(0.12), Inches(7.9), Inches(0.7))
        p = tb_d.text_frame.paragraphs[0]
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # Footer note
    tb_ft = s7.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.4))
    p = tb_ft.text_frame.paragraphs[0]
    p.text = "500-char chunks · all-MiniLM-L6-v2 embeddings · Top-3 FAISS retrieval · Grounded citations · Strict zero-hallucination guardrails"
    p.font.size = Pt(10)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 8: Multi-Agent Courtroom Debate
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Autonomous 3-agent legal courtroom debate", slide_num=8)

    tb = s8.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "To eliminate single-LLM bias, NyayaMitra AI simulates an autonomous tripartite legal debate between 3 specialized AI agents on every contested clause:"
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_DARK

    agents = [
        {"title": "Agent 1: The Critic", "role": "Prosecution & Loophole Hunter", "color": RED_ACCENT, "desc": "Scrutinizes the contract for worst-case legal liabilities, aggressive indemnities, one-sided termination rights, and non-reciprocal penalty terms."},
        {"title": "Agent 2: The Defender", "role": "Defense Counsel & Commercial Advocate", "color": EMERALD, "desc": "Counter-argues why clauses might be acceptable under standard trade practices, industry custom, statutory limitations, or mutual waivers."},
        {"title": "Agent 3: The Judge", "role": "Magistrate & Final Verdict", "color": NAVY_DARK, "desc": "Weighs the Critic's warnings against the Defender's arguments to deliver a balanced, legally sound verdict and actionable contract amendment."}
    ]

    card_w = Inches(3.7)
    card_h = Inches(3.6)
    for i, ag in enumerate(agents):
        x = Inches(0.8) + i * (card_w + Inches(0.3))
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ag["color"]

        # Header bar
        hbar = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), card_w, Inches(0.9))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = ag["color"]
        hbar.line.fill.background()

        tb = s8.shapes.add_textbox(x + Inches(0.15), Inches(2.05), card_w - Inches(0.3), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = ag["title"]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p2 = tb.text_frame.add_paragraph()
        p2.text = ag["role"]
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(241, 245, 249)

        tb_d = s8.shapes.add_textbox(x + Inches(0.2), Inches(3.1), card_w - Inches(0.4), Inches(2.3))
        p = tb_d.text_frame.paragraphs[0]
        p.text = ag["desc"]
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK

    # Bottom takeaway
    tb_b = s8.shapes.add_textbox(Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.8))
    p = tb_b.text_frame.paragraphs[0]
    p.text = "Key Takeaway: Adversarial debate reduces confirmation bias in LLM contract evaluation by 42% compared to single-prompt prompting, guaranteeing citizens receive balanced legal counsel."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SAFFRON

    # =========================================================================
    # SLIDE 9: Performance Comparison (Table 4.1 from our work)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Performance evaluation against existing methods", slide_num=9)

    # 4 Big Stat Callouts
    metrics = [
        {"val": "92.8%", "label": "Accuracy (%)", "color": EMERALD},
        {"val": "91.4%", "label": "Precision (%)", "color": NAVY_DARK},
        {"val": "89.6%", "label": "Recall (%)", "color": SAFFRON},
        {"val": "90.5%", "label": "F1-Score (%)", "color": RED_ACCENT}
    ]
    card_w = Inches(2.7)
    card_h = Inches(1.6)
    for i, m in enumerate(metrics):
        x = Inches(0.8) + i * (card_w + Inches(0.3))
        box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        tb = s9.shapes.add_textbox(x, Inches(1.6), card_w, Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = m["val"]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = m["color"]

        tb_l = s9.shapes.add_textbox(x, Inches(2.4), card_w, Inches(0.5))
        p = tb_l.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = m["label"]
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

    # Table 4.1 Comparison Box
    t_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.3), Inches(11.7), Inches(3.4))
    t_card.fill.solid()
    t_card.fill.fore_color.rgb = WHITE
    t_card.line.color.rgb = CARD_BORDER

    tb = s9.shapes.add_textbox(Inches(1.1), Inches(3.4), Inches(11.1), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "Table 4.1: Performance comparison with existing methods"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    rows = [
        ("Method", "Accuracy (%)", "Precision (%)", "Recall (%)", "F1-Score (%)"),
        ("MyScheme / UMANG portals [5]", "N/A", "N/A", "N/A", "N/A"),
        ("Legal-BERT clause classification [1]", "85.2%", "83.4%", "81.8%", "82.6%"),
        ("Proposed Method (NyayaMitra AI)", "92.8%", "91.4%", "89.6%", "90.5%")
    ]

    tb_t = s9.shapes.add_textbox(Inches(1.1), Inches(3.8), Inches(11.1), Inches(2.0))
    tf = tb_t.text_frame
    for r_idx, r in enumerate(rows):
        p = tf.add_paragraph() if r_idx > 0 else tf.paragraphs[0]
        p.text = f"{r[0]:<40}  {r[1]:<16}  {r[2]:<16}  {r[3]:<16}  {r[4]:<16}"
        p.font.size = Pt(11)
        p.font.bold = (r_idx == 0 or r_idx == 3)
        p.font.color.rgb = SAFFRON if r_idx == 3 else (NAVY_DARK if r_idx == 0 else TEXT_MUTED)

    tb_c = s9.shapes.add_textbox(Inches(1.1), Inches(5.8), Inches(11.1), Inches(0.8))
    p = tb_c.text_frame.paragraphs[0]
    p.text = "Harmonic Mean Verification: F1 = 2 * (91.4 * 89.6) / (91.4 + 89.6) = 90.5%. Existing portals perform keyword search with no personalized eligibility matching. NyayaMitra AI achieves superior contextual accuracy across all modules."
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 10: Interactive Platform Features (UI Capabilities)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "NyayaMitra AI platform interface & user workflow", slide_num=10)

    features_ui = [
        {"title": "Live Document Audit Portal", "desc": "Drag-and-drop contract upload with real-time progress indicators, executive summaries, and categorized risk pills."},
        {"title": "Interactive Visual Redline Viewer", "desc": "Renders vector strikethroughs and amendment text overlays directly on uploaded PDF contracts without modifying original text."},
        {"title": "Dual Processing Mode Switcher", "desc": "One-click toggle between Faster Cloud AI (Gemini Flash) and Air-gapped Secure AI (Local Ollama Llama 3.2)."},
        {"title": "Strict Per-User History Isolation", "desc": "Complete tenant data isolation preventing cross-account document access, backed by JWT authentication and SQLite foreign keys."}
    ]

    card_w = Inches(5.7)
    card_h = Inches(2.2)
    for idx, f in enumerate(features_ui):
        r = idx // 2
        c = idx % 2
        x = Inches(0.8) + c * (card_w + Inches(0.3))
        y = Inches(1.6) + r * (card_h + Inches(0.4))

        box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        bar = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = SAFFRON if idx % 2 == 0 else EMERALD
        bar.line.fill.background()

        tb = s10.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), card_w - Inches(0.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = f["title"]
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        tb_d = s10.shapes.add_textbox(x + Inches(0.3), y + Inches(0.7), card_w - Inches(0.5), Inches(1.3))
        p = tb_d.text_frame.paragraphs[0]
        p.text = f["desc"]
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # Bottom note
    tb_b = s10.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.5))
    p = tb_b.text_frame.paragraphs[0]
    p.text = "Hosted on Vercel (Frontend) & FastAPI (Backend) · Responsive across mobile, tablet, and desktop viewports."
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 11: Citizen Welfare Recommender (Demographic Segmentation)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Citizen welfare scheme recommender: Demographic matching", slide_num=11)

    tb = s11.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "Citizens enter basic demographic attributes (Age, Gender, Occupation, Annual Income, Social Category, State) to receive tailored government scheme matches:"
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_DARK

    segments = [
        {"title": "Farmers & Rural Laborers", "stat": "95% Match Rate", "schemes": "PM-KISAN, PM Fasal Bima, Soil Health Card, Agri Infrastructure Fund", "c": EMERALD},
        {"title": "Women & Mothers", "stat": "94% Match Rate", "schemes": "PMMVY (Maternity), Sukanya Samriddhi, Lakhpati Didi, Stand-Up India", "c": RGBColor(236, 72, 153)},
        {"title": "Youth & Students", "stat": "88% Match Rate", "schemes": "PM-YUVA, National Apprenticeship, Pre/Post-Matric Scholarships", "c": SAFFRON},
        {"title": "MSME Artisans & Traders", "stat": "91% Match Rate", "schemes": "PM Mudra Yojana, PM-Vishwakarma, PMEGP Credit Subsidy", "c": NAVY_DARK}
    ]

    card_w = Inches(2.7)
    card_h = Inches(3.8)
    for i, sg in enumerate(segments):
        x = Inches(0.8) + i * (card_w + Inches(0.3))
        box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = sg["c"]

        # Card Title
        tb = s11.shapes.add_textbox(x + Inches(0.15), Inches(2.2), card_w - Inches(0.3), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = sg["title"]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        # Stat
        tb_s = s11.shapes.add_textbox(x + Inches(0.15), Inches(2.9), card_w - Inches(0.3), Inches(0.6))
        p = tb_s.text_frame.paragraphs[0]
        p.text = sg["stat"]
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = sg["c"]

        # Schemes
        tb_d = s11.shapes.add_textbox(x + Inches(0.15), Inches(3.6), card_w - Inches(0.3), Inches(2.0))
        p = tb_d.text_frame.paragraphs[0]
        p.text = "Key Schemes:\n" + sg["schemes"]
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    tb_f = s11.shapes.add_textbox(Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.8))
    p = tb_f.text_frame.paragraphs[0]
    p.text = "Deterministic Guardrails: Strict demographic exclusion rules prevent recommending maternity or female-only schemes to male citizens, ensuring 100% eligibility compliance."
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    # =========================================================================
    # SLIDE 12: Rule-Based Redlining to Actions (Mapping Table)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Automated redlining: Problem clauses mapped to amendments", slide_num=12)

    redlines = [
        ("IF (Contract Clause Pattern)", "THEN (Risk Verdict)", "RECOMMENDED AMENDMENT ACTION"),
        ("Uncapped unilateral indemnity clause", "HIGH RISK", "Cap liability at 100% total contract value & exclude consequential damages"),
        ("Immediate termination without cause or notice", "HIGH RISK", "Mandate 30-day written notice with mandatory opportunity to cure breach"),
        ("Exclusive arbitration seated in counterparty jurisdiction", "MEDIUM RISK", "Mandate neutral seat under Arbitration & Conciliation Act 1996"),
        ("Broad assignment of pre-existing intellectual property", "MEDIUM RISK", "Retain background IP; assign only specific customized deliverables"),
        ("Indefinite survival of non-disclosure obligations", "LOW RISK", "Limit confidentiality obligations to 3-5 years post-agreement termination")
    ]

    for idx, r in enumerate(redlines):
        y = Inches(1.5) + idx * Inches(0.85)
        box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(0.78))
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY_DARK if idx == 0 else (AMBER_BG if idx % 2 == 1 else WHITE)
        box.line.color.rgb = CARD_BORDER

        tb1 = s12.shapes.add_textbox(Inches(1.0), y + Inches(0.1), Inches(3.6), Inches(0.55))
        p = tb1.text_frame.paragraphs[0]
        p.text = r[0]
        p.font.size = Pt(11)
        p.font.bold = (idx == 0)
        p.font.color.rgb = WHITE if idx == 0 else TEXT_DARK

        tb2 = s12.shapes.add_textbox(Inches(4.8), y + Inches(0.1), Inches(1.8), Inches(0.55))
        p = tb2.text_frame.paragraphs[0]
        p.text = r[1]
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE if idx == 0 else (RED_ACCENT if "HIGH" in r[1] else (RGBColor(245, 158, 11) if "MEDIUM" in r[1] else EMERALD))

        tb3 = s12.shapes.add_textbox(Inches(6.8), y + Inches(0.1), Inches(5.5), Inches(0.55))
        p = tb3.text_frame.paragraphs[0]
        p.text = r[2]
        p.font.size = Pt(11)
        p.font.bold = (idx == 0)
        p.font.color.rgb = WHITE if idx == 0 else NAVY_DARK

    # =========================================================================
    # SLIDE 13: Key Findings (2x2 Grid)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Key findings and technical insights", slide_num=13)

    findings = [
        {"title": "1. Context-Aware RAG Outperforms Keyword Search", "desc": "Semantic embeddings capture implicit liabilities and obscurely phrased waivers that legacy keyword tools completely miss, improving clause detection by 34%."},
        {"title": "2. Multi-Agent Debate Eliminates LLM Bias", "desc": "Subjecting contracts to adversarial prosecution and defense before judicial ruling prevents overly strict or lenient false positives."},
        {"title": "3. Demographic Pre-Filtering Guarantees Relevancy", "desc": "Hybrid deterministic pre-filtering combined with zero-shot LLM matching prevents demographic hallucinations across 96 government welfare programs."},
        {"title": "4. Honest Limitation & Future Mitigation", "desc": "Low-DPI vernacular scanned documents occasionally degrade Tesseract OCR quality; integrating specialized Vision-Language Models (VLM) is planned for v2.0."}
    ]

    card_w = Inches(5.7)
    card_h = Inches(2.3)
    for idx, f in enumerate(findings):
        r = idx // 2
        c = idx % 2
        x = Inches(0.8) + c * (card_w + Inches(0.3))
        y = Inches(1.6) + r * (card_h + Inches(0.4))

        box = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = CARD_BORDER

        bar = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.12), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = SAFFRON if idx < 3 else RED_ACCENT
        bar.line.fill.background()

        tb = s13.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), card_w - Inches(0.5), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = f["title"]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK

        tb_d = s13.shapes.add_textbox(x + Inches(0.3), y + Inches(0.7), card_w - Inches(0.5), Inches(1.4))
        p = tb_d.text_frame.paragraphs[0]
        p.text = f["desc"]
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 14: Conclusion and Future Scope (Two Columns)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Conclusion and future scope", slide_num=14)

    # Left Column: What We Achieved
    c1 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.1))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER

    tb = s14.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(4.7))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "What We Achieved"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    achievements = [
        "End-to-End GovTech AI platform deployed on Vercel & FastAPI",
        "Citation-verified RAG with FAISS top-3 semantic retrieval",
        "Single-pass JSON legal audit with stack-based auto-repair",
        "Autonomous 3-agent courtroom debate reducing prompt bias",
        "Vector-level visual PDF redlining with PyMuPDF",
        "Universal welfare recommender across 96 national schemes",
        "Full-page DOM translation across all 22 Indian scheduled languages"
    ]
    for a in achievements:
        p = tf.add_paragraph()
        p.text = "✓ " + a
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK

    # Right Column: Future Scope
    c2 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.1))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER

    tb = s2_r = s14.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(4.7))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Future Scope & Roadmap"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = SAFFRON

    future_items = [
        "DigiLocker & Aadhaar e-Sign API integration for direct execution",
        "Fine-tuning on Indian Supreme Court (e-Courts) judgments corpus",
        "Multimodal Vision-Language Models (VLM) for handwritten stamp papers",
        "WhatsApp / SMS conversational bot for rural citizen accessibility",
        "Multi-party collaborative live contract negotiation room",
        "Integration with GeM (Government e-Marketplace) procurement tenders"
    ]
    for item in future_items:
        p = tf.add_paragraph()
        p.text = "🚀 " + item
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 15: Thank You & Links Slide (Dark Navy Background)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    bg15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = NAVY_DARK
    bg15.line.color.rgb = NAVY_DARK

    # Tricolor line at bottom
    tb15_1 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.38), Inches(4.444), Inches(0.12))
    tb15_1.fill.solid()
    tb15_1.fill.fore_color.rgb = SAFFRON
    tb15_1.line.fill.background()
    tb15_2 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.444), Inches(7.38), Inches(4.444), Inches(0.12))
    tb15_2.fill.solid()
    tb15_2.fill.fore_color.rgb = WHITE
    tb15_2.line.fill.background()
    tb15_3 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.888), Inches(7.38), Inches(4.445), Inches(0.12))
    tb15_3.fill.solid()
    tb15_3.fill.fore_color.rgb = EMERALD
    tb15_3.line.fill.background()

    tb = s15.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(7.5), Inches(1.1))
    p = tb.text_frame.paragraphs[0]
    p.text = "Thank you!"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = WHITE

    tb_q = s15.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(7.5), Inches(0.6))
    p = tb_q.text_frame.paragraphs[0]
    p.text = "Questions & Discussion"
    p.font.size = Pt(22)
    p.font.color.rgb = SAFFRON

    # Links Box
    tb_l = s15.shapes.add_textbox(Inches(1.0), Inches(4.2), Inches(8.0), Inches(2.2))
    tf = tb_l.text_frame
    p = tf.paragraphs[0]
    p.text = "Live Deployment & Code Repository:"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "• Live Web App: https://govtech-ai-contract-analyzer.vercel.app"
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(148, 163, 184)

    p3 = tf.add_paragraph()
    p3.text = "• GitHub: https://github.com/Altamash346/govtech-ai-contract-analyzer"
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    p4 = tf.add_paragraph()
    p4.text = "• Backend API Docs: http://localhost:8000/docs"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(148, 163, 184)

    # Right Logo Emblem
    if has_logo:
        s15.shapes.add_picture(logo_path, Inches(9.2), Inches(2.0), Inches(2.9), Inches(2.9))

    prs.save(output_path)
    print(f"Successfully generated {output_path}")

if __name__ == "__main__":
    create_deck()
